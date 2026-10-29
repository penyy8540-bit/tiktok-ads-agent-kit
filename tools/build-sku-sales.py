#!/usr/bin/env python
"""Extract the operator's per-SKU daily sales/profit data into normalized JSON.

Source: their own `sku 销量情况*.xlsx` workbook.
  * `数据源` sheet is authoritative for daily numbers. It holds one or two
    side-by-side blocks, each block = 销售日期 | Seller SKU | 含税销售额 | 销量 |
    不含税销售额 | 毛利 | 毛利率.
  * per-station sheets (`de`, `uk`, ...) provide SKU → 产品名称/定价/毛利率 and
    the 7A-7D week buckets.

Usage:
  python build-sku-sales.py <workbook.xlsx> <out.json> [--dates 2026-10-24,...]

Output: {
  generatedAt, source, availableDates, dates, stations,
  rows: [{station, sku, name, price, grossMargin,
          daily: {"YYYY-MM-DD": {qty, revenueTaxIncl, revenueTaxExcl, profit, margin}},
          totals: {qty, revenueTaxIncl, revenueTaxExcl, profit}}]
}
"""
import datetime
import json
import os
import sys

SOURCE_SHEET = "数据源"
BLOCK_WIDTH = 7


def as_date(value):
    if isinstance(value, (datetime.datetime, datetime.date)):
        return value.strftime("%Y-%m-%d")
    text = str(value or "").strip()
    if len(text) >= 10 and text[4] == "-" and text[7] == "-":
        return text[:10]
    return None


def to_number(value):
    if value is None or value == "":
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).replace(",", "").replace("%", "").strip()
    try:
        return float(text)
    except ValueError:
        return 0.0


def station_of(sku):
    upper = sku.upper()
    for suffix, station in (("DE", "DE"), ("UK", "UK"), ("FR", "FR"), ("IT", "IT"), ("ES", "ES"), ("US", "US")):
        if upper.endswith(suffix):
            return station
    return ""


def read_meta(wb):
    """sku -> {name, price, margin, station} from the per-station sheets."""
    meta = {}
    for ws in wb.worksheets:
        if ws.title in {SOURCE_SHEET, "Instruction", "HiddenStyle", "HiddenAttr", "Template", "TemplateConfig", "Brand"}:
            continue
        grid = [list(r) for r in ws.iter_rows(values_only=True)]
        header_index = -1
        for index, row in enumerate(grid[:10]):
            if row and str(row[0] or "").strip().upper() == "SKU":
                header_index = index
                break
        if header_index < 0:
            continue
        for row in grid[header_index + 1:]:
            sku = str(row[0] or "").strip() if row else ""
            if not sku or sku in {"SKU", "合计", "总计"}:
                continue
            meta[sku] = {
                "name": str(row[1] or "").strip() if len(row) > 1 else "",
                "price": to_number(row[2] if len(row) > 2 else None),
                "margin": to_number(row[4] if len(row) > 4 else None),
                "station": ws.title.strip().upper(),
            }
    return meta


def read_daily(wb):
    """(sku, date) -> metrics, from every block of the 数据源 sheet.

    Blocks are found by scanning for date cells (the sheet places two blocks
    side by side with a spacer column between them, so a fixed stride is unsafe).
    """
    if SOURCE_SHEET not in wb.sheetnames:
        return {}, []
    ws = wb[SOURCE_SHEET]
    grid = [list(r) for r in ws.iter_rows(values_only=True)]
    daily = {}
    dates = set()
    for row in grid:
        if not row:
            continue
        index = 0
        while index < len(row):
            date = as_date(row[index])
            if date is None:
                index += 1
                continue
            if len(row) <= index + 3:
                break
            sku = str(row[index + 1] or "").strip()
            if not sku or sku.upper() in {"SELLER SKU", "SKU"}:
                index += 1
                continue
            metrics = {
                "revenueTaxIncl": to_number(row[index + 2] if len(row) > index + 2 else None),
                "qty": to_number(row[index + 3] if len(row) > index + 3 else None),
                "revenueTaxExcl": to_number(row[index + 4] if len(row) > index + 4 else None),
                "profit": to_number(row[index + 5] if len(row) > index + 5 else None),
                "margin": to_number(row[index + 6] if len(row) > index + 6 else None),
            }
            dates.add(date)
            bucket = daily.setdefault(sku, {})
            existing = bucket.get(date)
            if existing is None:
                bucket[date] = metrics
            else:
                for key, value in metrics.items():
                    existing[key] = existing.get(key, 0) + value
            index += BLOCK_WIDTH
    return daily, sorted(dates)


def main():
    if len(sys.argv) < 3:
        print("usage: build-sku-sales.py <workbook.xlsx> <out.json> [--dates a,b]", file=sys.stderr)
        return 2
    source, target = sys.argv[1], sys.argv[2]
    wanted = []
    if "--dates" in sys.argv:
        try:
            wanted = [d.strip() for d in sys.argv[sys.argv.index("--dates") + 1].split(",") if d.strip()]
        except IndexError:
            wanted = []
    if not os.path.exists(source):
        print(f"找不到文件：{source}", file=sys.stderr)
        return 1

    from openpyxl import load_workbook

    wb = load_workbook(source, read_only=True, data_only=True)
    try:
        meta = read_meta(wb)
        daily, available_dates = read_daily(wb)
        stations = sorted({entry["station"] for entry in meta.values() if entry.get("station")})
    finally:
        wb.close()

    keep_dates = [d for d in wanted if d in available_dates] if wanted else available_dates
    rows = []
    for sku, per_date in daily.items():
        info = meta.get(sku, {"name": "", "price": 0.0, "margin": 0.0, "station": station_of(sku)})
        selected = {d: per_date[d] for d in keep_dates if d in per_date}
        totals = {
            "qty": sum(v.get("qty", 0) for v in selected.values()),
            "revenueTaxIncl": round(sum(v.get("revenueTaxIncl", 0) for v in selected.values()), 2),
            "revenueTaxExcl": round(sum(v.get("revenueTaxExcl", 0) for v in selected.values()), 2),
            "profit": round(sum(v.get("profit", 0) for v in selected.values()), 2),
        }
        rows.append(
            {
                "station": info.get("station") or station_of(sku),
                "sku": sku,
                "name": info.get("name", ""),
                "price": info.get("price", 0.0),
                "grossMargin": info.get("margin", 0.0),
                "daily": selected,
                "totals": totals,
            }
        )
    rows.sort(key=lambda row: (-row["totals"]["qty"], row["station"], row["sku"]))

    payload = {
        "generatedAt": datetime.datetime.now().isoformat(timespec="seconds"),
        "source": source,
        "availableDates": available_dates,
        "dates": keep_dates,
        "stations": stations,
        "rows": rows,
    }
    with open(target, "w", encoding="utf-8") as handle:
        json.dump(payload, handle, ensure_ascii=False, indent=2)
    print(f"rows={len(rows)} dates={keep_dates} available={available_dates[-3:]}", file=sys.stderr)
    return 0


if __name__ == "__main__":
    sys.exit(main())
